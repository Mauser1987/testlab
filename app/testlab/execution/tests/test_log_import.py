"""Manual-log import from Excel (M2F): reading, mapping, interpretation, duplicates, events, commit, undo.

reference/log_import/TCN01.xlsx (a real legacy log, kept outside the repository) is the first validation
fixture; those tests are skipped when it is not present. Everything else uses small workbooks built here.
"""

import hashlib
import io
import zipfile
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from django.utils import timezone

from testlab.execution import manual_log
from testlab.execution.log_import import interpret, mapping, service, xlsx
from testlab.execution.models import ManualLogEntry, ManualLogImport, ManualLogLayoutMap, ManualMeasurement
from testlab.results import services as result_services
from testlab.results.models import FinalResultError, TestResult
from testlab.results.tests.helpers import log, make_run

pytestmark = pytest.mark.django_db

TCN01 = Path(settings.REPO_DIR) / "reference" / "log_import" / "TCN01.xlsx"
needs_tcn01 = pytest.mark.skipif(not TCN01.exists(), reason="reference/log_import/TCN01.xlsx not present")

HEADER = ["Date", "Time", "By", "RPM", "Torque\n[Nm]", "Plug Leakage\n[ml]", "Over\n[hours]",
          "Test Leakage\n[ml]", "Over\n[hours]", "Notes"]  # fmt: skip


def workbook(rows, header=HEADER, groups=None, meta=(), sheet="Log"):
    """An .xlsx: meta rows ("label:", value), an optional group row, the header, then the data rows."""
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = sheet
    number = 1
    for label, value in meta:
        ws.cell(number, 1, label)
        ws.cell(number, 3, value)
        number += 1
    if groups:
        for index, group in enumerate(groups, 1):
            if group:
                ws.cell(number, index, group)
        number += 1
    for index, heading in enumerate(header, 1):
        ws.cell(number, index, heading)
    for row in rows:
        number += 1
        for index, value in enumerate(row, 1):
            if value is not None:
                ws.cell(number, index, value)
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def reading(day, clock, by="JD", rpm=1100, torque=0.7, plug=None, plug_over=None, test=None, test_over=None,
            note=None):  # fmt: skip
    return [datetime(2026, 3, day), clock, by, rpm, torque, plug, plug_over, test, test_over, note]


@pytest.fixture
def run(user):
    run = make_run(user, "single", plug_seal=True, project_no="2025RSD02TL",
                   test=("C0625623AP", "C0625624SL"))  # fmt: skip
    setup = run.current_revision().setup
    setup.test_unit, setup.speed, setup.process_pressure = "B", "1100", "80"
    setup.save()
    return run


def upload(run, user, content, name="log.xlsx"):
    return service.upload(run, SimpleUploadedFile(name, content), user)


def confirm(batch, user, changes=None, save_layout=True):
    """Confirm the proposals (Review required ticked), with `changes` {letter: target}."""
    chosen = {m["letter"]: m["target"] for m in batch.mapping}
    chosen.update(changes or {})
    reviewed = {m["letter"] for m in batch.mapping if m["status"] == mapping.REVIEW}
    return service.confirm_mapping(batch, chosen, reviewed, user, save_layout=save_layout)


def staged(run, user, rows, **kwargs):
    return confirm(upload(run, user, workbook(rows, **kwargs)), user)


def rows_of(batch):
    return {row.sheet_row: row for row in batch.rows.all()}


def tcn01(run, user):
    return confirm(upload(run, user, TCN01.read_bytes(), "TCN01.xlsx"), user)


def make_final(run, user):
    result = result_services.create_result(run, user)
    TestResult.objects.filter(pk=result.pk).update(  # bypasses finalization: only the status matters here
        status=TestResult.Status.FINAL, finalized_at=timezone.now(), finalized_by=user)


# --- reading the workbook ------------------------------------------------------------------------


@needs_tcn01
def test_tcn01_two_row_header_rows_and_trailing_formula_rows():
    sheet = xlsx.read_sheet(TCN01.read_bytes())
    assert (sheet.name, sheet.header_row, sheet.group_row) == ("TCN001", 7, 6)
    assert [r for r, _ in sheet.rows] == list(range(8, 29))
    # the prepared rows below the log hold only formulas: ignored and listed (completely empty rows are skipped)
    assert len(sheet.ignored_rows) == 75 and min(sheet.ignored_rows) == 29 and max(sheet.ignored_rows) == 106
    columns = {c.letter: c for c in sheet.columns}
    assert (columns["J"].group, columns["J"].label, columns["J"].unit) == ("Job", "Job FI Temp.", "°C")
    assert columns["Q"].group == "Job" and columns["T"].group == "PLUG"  # carried to the columns with a unit
    assert columns["V"].label == "Notes:" and columns["V"].unit == ""


@needs_tcn01
def test_tcn01_formula_cells_keep_formula_and_stored_result():
    rows = dict(xlsx.read_sheet(TCN01.read_bytes()).rows)
    running = rows[11]["C"]
    assert running["f"] == '=IF(A11="","",((A11+B11)-($A$10+$B$10))*24)' and running["k"] == "number"
    assert rows[10]["Q"] == {"k": "number", "v": "0", "raw": "0"}  # typed 0, no formula


def with_stored_number(content, written, stored):
    """The workbook with the number text `written` replaced by `stored` in the sheet XML (as Excel may store it)."""
    source, target = zipfile.ZipFile(io.BytesIO(content)), io.BytesIO()
    with zipfile.ZipFile(target, "w") as patched:
        for item in source.infolist():
            data = source.read(item)
            if item.filename.startswith("xl/worksheets/"):
                data = data.replace(f"<v>{written}</v>".encode(), f"<v>{stored}</v>".encode())
            patched.writestr(item, data)
    return target.getvalue()


def test_numbers_are_read_from_the_file_text_to_excel_precision():
    content = with_stored_number(workbook([reading(2, time(8, 0), torque=0.75, plug=80.31)]),
                                 "0.75", "0.30000000000000004")  # fmt: skip
    cells = dict(xlsx.read_sheet(content).rows)[2]
    assert cells["E"]["raw"] == "0.30000000000000004" and cells["F"]["raw"] == "80.31"
    assert interpret.excel_decimal(cells["E"]["raw"]) == Decimal("0.3")
    assert interpret.excel_decimal("80.31") == Decimal("80.31") and str(interpret.excel_decimal("1100")) == "1.1E+3"


def test_legacy_xls_and_other_files_are_refused(run, user):
    with pytest.raises(service.LogImportError, match="save it as .xlsx"):
        upload(run, user, b"\xd0\xcf\x11\xe0" + b"\0" * 100, "old.xls")
    with pytest.raises(service.LogImportError, match="not an .xlsx"):
        upload(run, user, b"Date;Time\n", "log.csv")
    with pytest.raises(service.LogImportError, match="No log table"):
        upload(run, user, workbook([], header=["a", "b"]))
    assert not ManualLogImport.objects.exists()


def test_upload_keeps_the_workbook_byte_for_byte(run, user):
    content = workbook([reading(2, time(8, 0))])
    batch = upload(run, user, content, "TCN 01.xlsx")
    with batch.raw_file.open("rb") as stored:
        assert stored.read() == content
    assert batch.sha256 == hashlib.sha256(content).hexdigest() and batch.size == len(content)
    assert batch.raw_file.name.startswith("projects/2025RSD02TL/runs/001/log_imports/")
    assert batch.original_name == "TCN 01.xlsx" and batch.uploaded_by == user and batch.status == "mapping"


# --- mapping -------------------------------------------------------------------------------------


@needs_tcn01
def test_tcn01_proposals_for_a_single_seal_with_plug_seal(run, user):
    batch = upload(run, user, TCN01.read_bytes())
    proposed = {m["letter"]: (m["target"], m["status"]) for m in batch.mapping}
    assert proposed == {
        "A": ("date", "proposed"), "B": ("time", "proposed"),
        "C": ("derived", "not_imported"), "D": ("derived", "not_imported"),
        "E": ("operator", "proposed"), "F": ("speed", "proposed"), "G": ("torque", "proposed"),
        # not assumed: the engineer confirms the circuits and "Seal Press"
        "H": ("process_pressure", "review"),
        "I": ("seal_support_flow", "review"), "J": ("seal_support_inlet_temp", "review"),
        "K": ("seal_support_outlet_temp", "review"),
        "L": ("process_flow", "review"), "M": ("process_inlet_temp", "review"),
        "N": ("process_outlet_temp", "review"),
        "O": ("plug_face_temp", "proposed"),
        "P": ("test_leakage", "proposed"), "Q": ("interval:test_leakage", "proposed"),
        "R": ("derived", "not_imported"),
        "S": ("plug_leakage", "proposed"), "T": ("interval:plug_leakage", "proposed"),
        "U": ("derived", "not_imported"), "V": ("note", "proposed"),
    }  # fmt: skip


@needs_tcn01
def test_review_required_columns_must_be_confirmed(run, user):
    batch = upload(run, user, TCN01.read_bytes())
    chosen = {m["letter"]: m["target"] for m in batch.mapping}
    with pytest.raises(service.LogImportError) as refused:
        service.confirm_mapping(batch, chosen, reviewed={"H", "I"}, user=user)
    assert sum("review required" in p for p in refused.value.problems) == 5  # J K L M N
    batch.refresh_from_db()
    assert batch.status == "mapping" and not batch.rows.exclude(classification="").exists()
    # changing a Review-required column to another field is a decision too
    chosen.update({"J": "ignore", "K": "ignore", "L": "ignore", "M": "ignore", "N": "ignore"})
    service.confirm_mapping(batch, chosen, reviewed={"H", "I"}, user=user)
    assert ManualLogImport.objects.get(pk=batch.pk).status == "preview"


def test_mapping_problems(run, user):
    batch = upload(run, user, workbook([reading(2, time(8, 0))]))
    chosen = {m["letter"]: m["target"] for m in batch.mapping}
    problems = mapping.validate(run, batch.columns, {**chosen, "E": "speed"}, set(), service.proposals_of(batch))
    assert any("more than one column (D, E)" in p for p in problems)
    problems = mapping.validate(run, batch.columns, {**chosen, "H": "inboard_leakage"}, set(), [])
    assert any("not a field of this run" in p for p in problems)  # a dual-seal code on a single seal
    problems = mapping.validate(run, batch.columns, {**chosen, "B": "ignore"}, set(), [])
    assert any("needs a Time column" in p for p in problems)
    problems = mapping.validate(run, batch.columns, {**chosen, "F": "ignore"}, set(), [])
    assert any("collection period for Plug Seal leakage" in p for p in problems)


def test_fingerprint_ignores_spacing_and_case_not_column_order():
    columns = [{"letter": "A", "group": "Job", "label": "Job FI Temp.", "unit": "°C"},
               {"letter": "B", "group": "", "label": "RPM", "unit": ""}]  # fmt: skip
    spaced = [{**columns[0], "label": "job   fi temp."}, columns[1]]
    swapped = [{**columns[1], "letter": "A"}, {**columns[0], "letter": "B"}]
    assert mapping.fingerprint(columns) == mapping.fingerprint(spaced) != mapping.fingerprint(swapped)


def test_confirmed_mapping_is_reused_for_the_same_layout_and_configuration(run, user, other_user):
    content = workbook([reading(2, time(8, 0), plug=5)], header=HEADER[:4] + ["Seal Press [barg]"] + HEADER[5:])
    first = upload(run, user, content)
    assert {m["letter"]: m["status"] for m in first.mapping}["E"] == "review"
    confirm(first, user, changes={"E": "process_pressure"})
    layout = ManualLogLayoutMap.objects.get()
    assert layout.reviewed_by == user and layout.source_batch == first

    other_run = make_run(other_user, "single", plug_seal=True, project_no="2026RSD09TL")
    second = upload(other_run, other_user, content)
    statuses = {m["letter"]: (m["target"], m["status"]) for m in second.mapping}
    assert statuses["E"] == ("process_pressure", "confirmed") and second.layout_map == layout
    # still reviewable: the engineer may choose another field
    confirm(second, other_user, changes={"E": "ignore"}, save_layout=False)
    assert ManualLogLayoutMap.objects.get().columns[4]["target"] == "process_pressure"

    dual = make_run(other_user, "dual_pressurized", project_no="2026RSD08TL")
    assert upload(dual, other_user, content).layout_map is None  # another configuration: proposed afresh


def test_each_batch_keeps_its_own_copy_of_the_mapping(run, user):
    content = workbook([reading(2, time(8, 0))])
    first = confirm(upload(run, user, content), user, changes={"C": "ignore"})
    confirm(upload(run, user, content), user)  # saves the layout again with the operator mapped
    first.refresh_from_db()
    assert {m["letter"]: m["target"] for m in first.mapping}["C"] == "ignore"


# --- interpretation ------------------------------------------------------------------------------


@needs_tcn01
def test_tcn01_values_zeros_blanks_and_operator(run, user):
    rows = rows_of(tcn01(run, user))
    start = rows[8].interpreted
    assert start["observed_at"] == "2025-06-30T08:00:00+02:00"
    assert start["values"]["speed"] == "0" and start["values"]["seal_support_flow"] == "0"  # zeros are readings
    assert "test_leakage" not in start["values"] and "plug_leakage" not in start["values"]  # blank: not measured
    assert rows[8].source_operator == "Martijn & Claude"
    assert rows[11].interpreted["values"]["test_leakage"] == "125" and rows[11].interpreted["leakage_hours"] is None
    # anomalous-looking values stay exactly as recorded, nothing is flagged
    assert rows[25].interpreted["values"]["process_outlet_temp"] == "31.95" and not rows[25].interpreted["flags"]
    # the typed 0 in Over [hours] at the start is not a collection period
    assert {f["cell"] for f in rows[10].interpreted["flags"]} == {"Q10", "T10"}
    assert rows[10].interpreted["leakage_hours"] is None


def test_an_unreadable_measurement_makes_the_whole_row_an_error(run, user):
    batch = staged(run, user, [reading(2, time(8, 0)), reading(2, time(9, 0), torque="n/a", plug=12)])
    row = rows_of(batch)[3]
    assert row.classification == "error" and row.decision == "skip"
    assert row.interpreted["errors"] == [{"cell": "E3", "message": "Torque [Nm]: “n/a” is not a number."}]
    service.save_choices(batch, {f"r{row.pk}_decision": "import"}, user)
    assert any("Row 3: error — it can only be skipped" in p for p in service.problems(batch))
    with pytest.raises(service.LogImportError):
        service.commit(batch, user)
    assert not ManualLogEntry.objects.exists()


def test_more_decimals_than_stored_is_an_error_never_rounded(run, user):
    row = rows_of(staged(run, user, [reading(2, time(8, 0), torque=0.12345)]))[2]
    assert row.classification == "error" and "more decimals" in row.interpreted["errors"][0]["message"]


def test_number_typed_as_text_is_read_as_the_entry_form_reads_it_and_flagged(run, user):
    row = rows_of(staged(run, user, [reading(2, time(8, 0), torque="0,75")]))[2]
    assert row.classification == "new" and row.interpreted["values"]["torque"] == "0.75"
    assert "Number stored as text" in row.interpreted["flags"][0]["message"]


def test_collection_periods(run, user):
    batch = staged(run, user, [
        reading(2, time(8, 0)),
        reading(2, time(9, 0), plug=10, plug_over="=A3-A2", test=5, test_over="=A3-A2"),  # formulas: default
        reading(2, time(10, 0), plug=10, plug_over=1.5, test=5, test_over=1.5),  # typed: imported
        reading(2, time(11, 0), plug=10, plug_over=0.5, test=5, test_over=1),  # disagree: choose
    ])  # fmt: skip
    rows = rows_of(batch)
    assert rows[3].interpreted["leakage_hours"] is None and not rows[3].interpreted["period_options"]
    assert rows[4].interpreted["leakage_hours"] == "1.5"
    assert rows[5].interpreted["period_options"] == ["0.5", "1", ""]
    assert any("Row 5: choose the collection period" in p for p in service.problems(batch))
    service.save_choices(batch, {f"r{rows[5].pk}_period": "0.5"}, user)
    service.commit(batch, user)
    periods = {timezone.localtime(e.observed_at).hour: e.leakage_hours for e in ManualLogEntry.objects.all()}
    assert periods == {8: None, 9: None, 10: Decimal("1.5"), 11: Decimal("0.5")}


def test_time_in_a_daylight_saving_change_is_an_error(run, user):
    row = rows_of(staged(run, user, [[datetime(2026, 10, 25), time(2, 30), "JD", 1100]]))[2]
    assert row.classification == "error" and "daylight-saving" in row.interpreted["errors"][0]["message"]


def test_rows_without_values_or_note_are_not_imported(run, user):
    row = rows_of(staged(run, user, [[datetime(2026, 3, 2), time(8, 0), "JD"]]))[2]
    assert row.classification == "empty" and row.decision == "skip"


# --- the rows above the log table ------------------------------------------------------------------


def test_what_is_written_above_the_log_table_is_not_read(run, user):
    meta = [("Test Project Number:", "2025RSD02TL"), ("Run no.:", "TCN002"), ("Test Unit:", "A"),
            ("Speed:", "1100 CW"), ("Stat Face:", "C0625623AP"), ("Data not Logged", None)]  # fmt: skip
    batch = staged(run, user, [reading(2, time(8, 0))], meta=meta)
    assert service.problems(batch) == []  # another run number and test unit above the table: not looked at
    assert service.commit(batch, user) == 1
    setup = manual_log.run_setup(run)
    assert (setup.test_unit, setup.speed, run.run_no) == ("B", "1100", 1)


# --- duplicates and conflicts --------------------------------------------------------------------


def test_new_duplicate_and_conflict_on_the_same_minute_only(run, user):
    log(run, user, 0, speed=1100, torque=Decimal("0.7"))  # START: 29-06-2026 09:00
    log(run, user, 1, speed=1100, torque=Decimal("0.8"))  # 10:00
    batch = staged(run, user, [
        [datetime(2026, 6, 29), time(9, 0, 40), "JD", 1100, 0.7],  # same minute, same values
        [datetime(2026, 6, 29), time(10, 0), "JD", 1100, 0.9],  # same minute, other torque
        [datetime(2026, 6, 29), time(10, 1), "JD", 1100, 0.8],  # next minute: a reading of its own
    ])  # fmt: skip
    rows = rows_of(batch)
    assert [(rows[n].classification, rows[n].decision) for n in (2, 3, 4)] == [
        ("duplicate", "skip"), ("conflict", ""), ("new", "import")]
    assert rows[3].match_details["differences"] == [
        {"code": "torque", "label": "Torque [Nm]", "existing": "0.8", "imported": "0.9"}]
    assert any("Row 3: conflicts" in p for p in service.problems(batch))


def test_a_conflict_imported_is_an_additional_entry_never_an_overwrite(run, user):
    existing = log(run, user, 0, speed=1100, torque=Decimal("0.7"), note="by hand")
    batch = staged(run, user, [[datetime(2026, 6, 29), time(9, 0), "JD", 1100, 0.9, None, None, None, None, "x"]])
    row = batch.rows.get()
    service.save_choices(batch, {f"r{row.pk}_decision": "import"}, user)
    service.commit(batch, user)
    existing.refresh_from_db()
    assert existing.note == "by hand" and existing.measurements.get(code="torque").value == Decimal("0.7")
    assert ManualLogEntry.objects.count() == 2


def test_rows_repeated_in_the_workbook(run, user):
    batch = staged(run, user, [reading(2, time(8, 0)), reading(2, time(8, 0)), reading(2, time(8, 0), rpm=0)])
    rows = rows_of(batch)
    assert [rows[n].classification for n in (2, 3, 4)] == ["new", "duplicate", "conflict"]
    assert rows[3].match_details["row"] == 2


@needs_tcn01
def test_importing_the_same_workbook_again_gives_duplicates_of_that_file(run, user):
    service.commit(tcn01(run, user), user)
    again = tcn01(run, user)
    rows = list(again.rows.all())
    assert {r.classification for r in rows} == {"duplicate"} and all(r.decision == "skip" for r in rows)
    assert all(r.match_details["same_file"] for r in rows)
    assert "imported from TCN01.xlsx row 8" in rows[0].match_details["source"]


# --- event suggestions ---------------------------------------------------------------------------


@needs_tcn01
def test_tcn01_event_suggestions(run, user):
    rows = rows_of(tcn01(run, user))
    assert rows[9].suggestions["category"]["value"] == "other"  # Start recording: not a shaft start
    assert rows[28].suggestions["category"]["value"] == "other"  # Stop recording
    assert rows[10].suggestions["category"]["value"] == "start"  # Start RPM
    timed = rows[11].suggestions["events"]
    assert [(e["time"], e["category"], e["needs_time"]) for e in timed] == [("09:36", "adjustment", False)]
    assert rows[11].suggestions["category"] is None  # the note describes the 09:36 event, not the reading
    after = rows[26].suggestions["events"]
    assert [(e["time"], e["category"], e["needs_time"]) for e in after] == [("", "stop", True)]
    assert "between 12:05 and 12:08" in after[0]["reason"] and rows[26].suggestions["category"] is None


@needs_tcn01
def test_suggestions_are_not_applied_unless_accepted(run, user):
    service.commit(tcn01(run, user), user)
    entries = list(ManualLogEntry.objects.order_by("observed_at"))
    assert len(entries) == 21 and not any(e.category for e in entries)
    assert entries[3].note == "9:36 Changed 1 coil 1 cooler Job and Plug  & Plug reduse cw in the cooler at 10L/min"


@needs_tcn01
def test_accepted_events_and_categories(run, user):
    batch = tcn01(run, user)
    rows = rows_of(batch)
    service.save_choices(batch, {
        f"r{rows[10].pk}_category": "start",
        f"r{rows[11].pk}_e0_accept": "1", f"r{rows[11].pk}_e0_time": "09:36", f"r{rows[11].pk}_e0_category": "adjustment",
        f"r{rows[26].pk}_e0_accept": "1", f"r{rows[26].pk}_e0_time": "", f"r{rows[26].pk}_e0_category": "stop",
    }, user)  # fmt: skip
    assert any("Row 26: enter the time" in p for p in service.problems(batch))
    data = {f"r{rows[10].pk}_category": "start", f"r{rows[11].pk}_e0_accept": "1", f"r{rows[11].pk}_e0_time": "09:36",
            f"r{rows[11].pk}_e0_category": "adjustment", f"r{rows[26].pk}_e0_accept": "1",
            f"r{rows[26].pk}_e0_time": "12:06", f"r{rows[26].pk}_e0_category": "stop"}  # fmt: skip
    service.save_choices(batch, data, user)
    assert service.commit(batch, user) == 23
    event = ManualLogEntry.objects.get(observed_at=datetime(2025, 6, 30, 9, 36, tzinfo=timezone.get_current_timezone()))
    assert event.category == "adjustment" and not event.measurements.exists() and event.source_row == rows[11]
    assert ManualLogEntry.objects.get(source_row=rows[10]).category == "start"
    stop = ManualLogEntry.objects.get(category="stop")
    assert timezone.localtime(stop.observed_at).strftime("%H:%M") == "12:06"
    # an event entry does not end a leakage collection: row 11's rate is still over the time since 09:23
    rows_by_time = {timezone.localtime(r.entry.observed_at).strftime("%d %H:%M"): r for r in manual_log.log_rows(run)}
    leak = rows_by_time["30 10:35"].leakage["test_leakage"]
    assert leak.hours == pytest.approx(1.2) and leak.rate == pytest.approx(125 / 1.2)


# --- commit ----------------------------------------------------------------------------------------


@needs_tcn01
def test_tcn01_import_creates_traceable_entries(run, user):
    batch = tcn01(run, user)
    assert service.problems(batch) == []
    assert service.commit(batch, user) == 21
    batch.refresh_from_db()
    assert batch.status == "imported" and batch.imported_by == user and batch.imported_at is not None
    entries = ManualLogEntry.objects.select_related("source_row").order_by("observed_at")
    first = entries[0]
    assert first.entered_by == user and first.source_row.sheet_row == 8
    assert first.source_row.source_operator == "Martijn & Claude" and first.note.startswith("Job 1 cooler")
    values = {m.code: m.value for m in first.measurements.all()}
    assert values["speed"] == 0 and values["torque"] == 0 and "test_leakage" not in values
    assert ManualMeasurement.objects.filter(code="process_outlet_temp", value=Decimal("31.95")).exists()
    assert all(row.outcome["entries"] for row in batch.rows.all())


def test_commit_is_all_or_nothing(run, user, monkeypatch):
    batch = staged(run, user, [reading(2, time(8, 0)), reading(2, time(9, 0))])
    created = []
    original = ManualMeasurement.objects.create

    def failing(**kwargs):
        if len(created) == 3:  # the second row's torque
            raise RuntimeError("disk full")
        created.append(kwargs)
        return original(**kwargs)

    monkeypatch.setattr(ManualMeasurement.objects, "create", failing)
    with pytest.raises(RuntimeError):
        service.commit(batch, user)
    assert not ManualLogEntry.objects.exists() and ManualLogImport.objects.get().status == "preview"


def test_a_log_changed_since_the_preview_stops_the_import(run, user):
    batch = staged(run, user, [[datetime(2026, 6, 29), time(9, 0), "JD", 1100, 0.9]])
    log(run, user, 0, speed=1100, torque=Decimal("0.7"))  # entered by hand in the meantime
    with pytest.raises(service.StaleImport, match="row 2 is now conflict"):
        service.commit(batch, user)
    row = batch.rows.get()
    assert row.classification == "conflict" and row.decision == "" and ManualLogEntry.objects.count() == 1


# --- undo ------------------------------------------------------------------------------------------


def test_undo_removes_only_the_batch_entries_and_keeps_the_record(run, user):
    manual = log(run, user, 0, speed=1100)
    batch = staged(run, user, [reading(2, time(8, 0)), reading(2, time(9, 0))])
    service.commit(batch, user)
    with pytest.raises(service.LogImportError, match="reason"):
        service.undo(batch, user, " ")
    assert service.undo(batch, user, "Wrong run") == 2
    batch.refresh_from_db()
    assert list(ManualLogEntry.objects.all()) == [manual]
    assert batch.status == "undone" and batch.undo_reason == "Wrong run" and batch.undone_by == user
    assert len(batch.undo_details["removed"]) == 2 and batch.rows.count() == 2
    assert batch.raw_file.storage.exists(batch.raw_file.name)
    with pytest.raises(service.LogImportError, match="only an imported batch"):
        service.undo(batch, user, "again")
    service.commit(staged(run, user, [reading(2, time(8, 0))]), user)  # the workbook can be imported again


def test_undo_asks_before_removing_entries_edited_after_the_import(run, user, other_user):
    batch = staged(run, user, [reading(2, time(8, 0))])
    service.commit(batch, user)
    entry = ManualLogEntry.objects.get()
    entry.note, entry.updated_by = "corrected", other_user
    entry.save()
    with pytest.raises(service.LogImportError, match="edited after the import"):
        service.undo(batch, user, "Wrong run")
    service.undo(batch, user, "Wrong run", include_edited=True)
    batch.refresh_from_db()
    assert batch.undo_details["edited"][0]["updated_by"] == "Mark Test" and not ManualLogEntry.objects.exists()


def test_undo_needs_the_permission(run, user, monkeypatch):
    batch = staged(run, user, [reading(2, time(8, 0))])
    service.commit(batch, user)
    monkeypatch.setattr(service, "can_undo_log_import", lambda u: False)
    with pytest.raises(service.LogImportError, match="not allowed"):
        service.undo(batch, user, "Wrong run")
    assert ManualLogEntry.objects.count() == 1


def test_discard_removes_a_batch_never_imported(run, user):
    batch = upload(run, user, workbook([reading(2, time(8, 0))]))
    name, storage = batch.raw_file.name, batch.raw_file.storage
    service.discard(batch)
    assert not ManualLogImport.objects.exists() and not storage.exists(name)


# --- final Result ----------------------------------------------------------------------------------


def test_a_final_result_closes_imports(run, user):
    imported = staged(run, user, [reading(2, time(8, 0))])
    service.commit(imported, user)
    open_batch = staged(run, user, [reading(2, time(9, 0))])
    make_final(run, user)
    with pytest.raises(service.LogImportError, match="is final"):
        upload(run, user, workbook([reading(2, time(10, 0))]))
    with pytest.raises(service.LogImportError, match="is final"):
        service.commit(open_batch, user)
    with pytest.raises(service.LogImportError, match="is final"):
        service.undo(imported, user, "Wrong run")
    with pytest.raises(service.LogImportError, match="is final"):
        service.discard(open_batch)
    with pytest.raises(FinalResultError):  # the model guard behind the pages
        imported.save()
    with pytest.raises(FinalResultError):
        imported.rows.first().save()
    assert ManualLogEntry.objects.count() == 1


# --- pages -----------------------------------------------------------------------------------------


def test_upload_mapping_preview_import_pages(client, run, user):
    client.force_login(user)
    url = reverse("execution:log_import", args=[run.pk])
    assert "Import Excel manual log" in client.get(url).content.decode()
    content = workbook([reading(2, time(8, 0), plug=5), reading(2, time(9, 0), plug=6, note="9:10 Changed valve")])
    response = client.post(url, {"file": SimpleUploadedFile("log.xlsx", content)})
    batch = ManualLogImport.objects.get()
    assert response.url == reverse("execution:log_import_mapping", args=[batch.pk])
    page = client.get(response.url).content.decode()
    assert "Column mapping" in page and "Plug Seal leakage [ml]" in page

    form = {f"target_{m['letter']}": m["target"] for m in batch.mapping} | {"save_layout": "1"}
    response = client.post(reverse("execution:log_import_mapping", args=[batch.pk]), form)
    assert response.url == reverse("execution:log_import_preview", args=[batch.pk])
    page = client.get(response.url).content.decode()
    assert "Detected event time 09:10 — create separate event?" in page and "JD" in page

    rows = {r.sheet_row: r for r in batch.rows.all()}
    decisions = {f"r{r.pk}_decision": "import" for r in rows.values()}
    response = client.post(reverse("execution:log_import_preview", args=[batch.pk]), decisions | {"action": "import"})
    assert response.url == reverse("execution:log", args=[run.pk]) and ManualLogEntry.objects.count() == 2
    log_page = client.get(response.url).content.decode()
    assert "JD" in log_page and reverse("execution:log_import_detail", args=[batch.pk]) in log_page
    record = client.get(reverse("execution:log_import_detail", args=[batch.pk])).content.decode()
    assert "Undo import" in record and "Mapping used" in record


def test_upload_page_explains_a_refused_file(client, run, user):
    client.force_login(user)
    response = client.post(reverse("execution:log_import", args=[run.pk]),
                            {"file": SimpleUploadedFile("old.xls", b"\xd0\xcf\x11\xe0" + b"\0" * 64)})  # fmt: skip
    assert response.status_code == 200 and "save it as .xlsx" in response.content.decode()


def test_undo_page(client, run, user):
    client.force_login(user)
    batch = staged(run, user, [reading(2, time(8, 0))])
    service.commit(batch, user)
    url = reverse("execution:log_import_undo", args=[batch.pk])
    assert "Reason" in client.get(url).content.decode()
    assert "Give the reason" in client.post(url, {"reason": ""}).content.decode()
    client.post(url, {"reason": "Imported into the wrong run"})
    assert ManualLogImport.objects.get().status == "undone" and not ManualLogEntry.objects.exists()
